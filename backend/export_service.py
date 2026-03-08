import io
import os
from datetime import datetime
from typing import Optional


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _c(v) -> str:
    try:
        return f"${float(v):,.2f}"
    except (TypeError, ValueError):
        return "–"


def _pct(v, total) -> str:
    try:
        if total and float(total) > 0:
            return f"{float(v) / float(total) * 100:.1f}%"
        return "0%"
    except (TypeError, ValueError):
        return "–"


def _f(v, decimals=2, suffix="") -> str:
    try:
        return f"{round(float(v), decimals):,}{suffix}"
    except (TypeError, ValueError):
        return str(v) if v is not None else "–"


def _pct_float(v, total) -> float:
    try:
        if total and float(total) > 0:
            return round(float(v) / float(total) * 100, 1)
        return 0.0
    except (TypeError, ValueError):
        return 0.0


# ─── Chart helpers (matplotlib → PNG bytes) ───────────────────────────────────

_CHART_COLORS = [
    "#2563EB", "#0891B2", "#7C3AED", "#D97706",
    "#059669", "#DC2626", "#EA580C", "#8B5CF6", "#10B981", "#F59E0B", "#EC4899",
]


def _make_bar_chart(labels, values, title="", colors=None, figsize=(7, 3.5)):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.ticker as mticker

    fig, ax = plt.subplots(figsize=figsize)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#F8FAFC")

    bar_colors = colors or _CHART_COLORS[:len(labels)]

    bars = ax.bar(labels, values, color=bar_colors, width=0.6, zorder=3)
    ax.set_title(title, color="#1E293B", fontsize=11, pad=10, fontweight="bold")
    ax.tick_params(colors="#475569", labelsize=8)
    ax.xaxis.set_tick_params(rotation=30)
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"${x:,.0f}"))
    for spine in ax.spines.values():
        spine.set_edgecolor("#CBD5E1")
    ax.yaxis.set_tick_params(colors="#475569")
    ax.xaxis.set_tick_params(colors="#475569")
    ax.grid(axis="y", color="#E2E8F0", linestyle="--", linewidth=0.7, zorder=0)
    ax.set_xlabel("")

    for bar, val in zip(bars, values):
        if val > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() * 1.02,
                f"${val:,.0f}",
                ha="center", va="bottom", color="#334155", fontsize=7, fontweight="bold"
            )

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _make_pie_chart(labels, values, title="", figsize=(5, 3.8)):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    filtered = [(l, v, c) for l, v, c in zip(labels, values, _CHART_COLORS[:len(labels)]) if v > 0]
    if not filtered:
        return None
    fl, fv, fc = zip(*filtered)

    fig, ax = plt.subplots(figsize=figsize)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    wedges, texts, autotexts = ax.pie(
        fv, labels=None, colors=fc, autopct="%1.1f%%",
        startangle=140, pctdistance=0.78,
        wedgeprops={"linewidth": 0.5, "edgecolor": "white"}
    )
    for at in autotexts:
        at.set_color("#1E293B")
        at.set_fontsize(7)
        at.set_fontweight("bold")

    ax.legend(
        wedges, [f"{l}" for l in fl],
        loc="center left", bbox_to_anchor=(1, 0.5),
        fontsize=8, frameon=False, labelcolor="#334155"
    )
    ax.set_title(title, color="#1E293B", fontsize=11, pad=8, fontweight="bold")

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _make_line_chart(quantity_curve, current_qty=None, figsize=(9, 3.8)):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.ticker as mticker

    if not quantity_curve:
        return None

    qs  = [pt["quantity"]     for pt in quantity_curve]
    cpp = [pt["cost_per_part"] for pt in quantity_curve]
    mat = [pt.get("material_cost", 0) for pt in quantity_curve]
    mac = [pt.get("machine_cost",  0) for pt in quantity_curve]

    fig, ax = plt.subplots(figsize=figsize)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#F8FAFC")

    ax.plot(qs, cpp, color="#2563EB", linewidth=2,   label="Total $/part")
    ax.plot(qs, mat, color="#0891B2", linewidth=1.4, label="Material $/part", linestyle="--")
    ax.plot(qs, mac, color="#7C3AED", linewidth=1.4, label="Machine $/part",  linestyle="--")

    if current_qty:
        ax.axvline(current_qty, color="#D97706", linestyle=":", linewidth=1.5,
                   label=f"Current qty ({current_qty:,})")

    ax.set_title("Cost per Part vs Production Quantity", color="#1E293B", fontsize=11, pad=10, fontweight="bold")
    ax.tick_params(colors="#475569", labelsize=8)
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"${x:,.0f}"))
    ax.xaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"{int(x):,}"))
    for spine in ax.spines.values():
        spine.set_edgecolor("#CBD5E1")
    ax.grid(color="#E2E8F0", linestyle="--", linewidth=0.6)
    ax.legend(fontsize=8, frameon=True, framealpha=0.8, edgecolor="#E2E8F0", labelcolor="#334155")

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


# ─── STL 3-D snapshot ─────────────────────────────────────────────────────────

def _render_stl_snapshot(
    file_path: str,
    build_direction: str = "z",
    overhang_angle: float = 45.0,
    figsize=(9.0, 4.2),
) -> Optional[bytes]:
    """Render STL as a 3-D matplotlib snapshot with overhang highlighting.

    Returns PNG bytes or None if the file is missing / rendering fails.
    """
    try:
        import sys, types, math

        # trimesh may need a networkx stub on minimal installs
        try:
            import networkx  # noqa: F401
        except ModuleNotFoundError:
            _nx = types.ModuleType("networkx")
            sys.modules["networkx"] = _nx

        import trimesh
        import numpy as np
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from mpl_toolkits.mplot3d.art3d import Poly3DCollection
        from matplotlib.patches import Patch

        if not os.path.exists(file_path):
            return None

        # ── Load mesh ─────────────────────────────────────────────────────────
        loaded = trimesh.load(file_path, force="mesh")
        if isinstance(loaded, trimesh.Scene):
            geoms = list(loaded.geometry.values())
            if not geoms:
                return None
            mesh = trimesh.util.concatenate(geoms) if len(geoms) > 1 else geoms[0]
        else:
            mesh = loaded

        if not isinstance(mesh, trimesh.Trimesh) or len(mesh.faces) == 0:
            return None

        # ── Overhang detection ────────────────────────────────────────────────
        axis_map = {"x": 0, "y": 1, "z": 2}
        build_axis = axis_map.get(build_direction.lower(), 2)
        build_vec = np.zeros(3)
        build_vec[build_axis] = 1.0

        verts   = np.asarray(mesh.vertices, dtype=float)
        faces   = np.asarray(mesh.faces,    dtype=int)
        normals = np.asarray(mesh.face_normals, dtype=float)
        dots    = normals @ build_vec
        threshold    = -math.cos(math.radians(overhang_angle))
        is_overhang  = dots < threshold

        normal_faces   = faces[~is_overhang]
        overhang_faces = faces[is_overhang]

        # Sub-sample for rendering performance (matplotlib painter's limit)
        RENDER_LIMIT = 4000
        rng = np.random.default_rng(42)

        def _sub(fa, n):
            if len(fa) <= n:
                return fa
            return fa[rng.choice(len(fa), n, replace=False)]

        normal_r   = _sub(normal_faces,   RENDER_LIMIT)
        overhang_r = _sub(overhang_faces, min(RENDER_LIMIT, max(1, len(overhang_faces))))

        # ── Bounding box ──────────────────────────────────────────────────────
        mins    = verts.min(axis=0)
        maxs    = verts.max(axis=0)
        extents = maxs - mins
        corners = np.array([
            [mins[0], mins[1], mins[2]], [maxs[0], mins[1], mins[2]],
            [maxs[0], maxs[1], mins[2]], [mins[0], maxs[1], mins[2]],
            [mins[0], mins[1], maxs[2]], [maxs[0], mins[1], maxs[2]],
            [maxs[0], maxs[1], maxs[2]], [mins[0], maxs[1], maxs[2]],
        ])
        bbox_edges = [
            (0,1),(1,2),(2,3),(3,0),
            (4,5),(5,6),(6,7),(7,4),
            (0,4),(1,5),(2,6),(3,7),
        ]

        # ── Figure ────────────────────────────────────────────────────────────
        fig = plt.figure(figsize=figsize, facecolor="white")
        ax  = fig.add_subplot(111, projection="3d")
        ax.set_facecolor("white")

        if len(normal_r) > 0:
            pc = Poly3DCollection(verts[normal_r], alpha=0.70, linewidths=0)
            pc.set_facecolor("#2563EB")
            ax.add_collection3d(pc)

        if len(overhang_r) > 0:
            pc2 = Poly3DCollection(verts[overhang_r], alpha=0.90, linewidths=0)
            pc2.set_facecolor("#DC2626")
            ax.add_collection3d(pc2)

        for a, b in bbox_edges:
            ax.plot3D(*zip(corners[a], corners[b]),
                      color="#D97706", linewidth=0.8, linestyle="--", alpha=0.85)

        # Build plate surface
        pm = max(extents[0], extents[1]) * 0.10
        xs = np.array([[mins[0]-pm, maxs[0]+pm], [mins[0]-pm, maxs[0]+pm]])
        ys = np.array([[mins[1]-pm, mins[1]-pm], [maxs[1]+pm, maxs[1]+pm]])
        zs = np.full_like(xs, mins[2])
        ax.plot_surface(xs, ys, zs, alpha=0.15, color="#0891B2", linewidth=0)

        # Build-direction arrow
        ctr   = (mins + maxs) / 2
        arrow = extents[build_axis] * 0.45 * build_vec
        ax.quiver(*ctr, *arrow, color="#059669", linewidth=1.8, arrow_length_ratio=0.22)

        # ── Axes / limits ─────────────────────────────────────────────────────
        pad = max(extents) * 0.04
        ax.set_xlim(mins[0] - pad, maxs[0] + pad)
        ax.set_ylim(mins[1] - pad, maxs[1] + pad)
        ax.set_zlim(mins[2],       maxs[2] + extents[build_axis] * 0.55)

        ax.tick_params(colors="#64748B", labelsize=6)
        ax.set_xlabel("X (mm)", color="#64748B", fontsize=6.5, labelpad=1)
        ax.set_ylabel("Y (mm)", color="#64748B", fontsize=6.5, labelpad=1)
        bax = "XYZ"[build_axis]
        ax.set_zlabel(f"{bax} / Build (mm)", color="#64748B", fontsize=6.5, labelpad=1)
        for pane in [ax.xaxis.pane, ax.yaxis.pane, ax.zaxis.pane]:
            pane.fill = False
            pane.set_edgecolor("#E2E8F0")

        n_overhang = int(is_overhang.sum())
        pct_overhang = n_overhang / max(len(faces), 1) * 100
        ax.set_title(
            f"Bounding box: {extents[0]:.1f} × {extents[1]:.1f} × {extents[2]:.1f} mm"
            f"   |   Build: +{bax}"
            f"   |   Overhang faces: {n_overhang} ({pct_overhang:.1f}%)",
            color="#334155", fontsize=7.5, pad=6,
        )

        legend_handles = [
            Patch(facecolor="#2563EB", label="Nominal surface"),
            Patch(facecolor="#DC2626", label=f"Overhang (>{overhang_angle:.0f}°)"),
            Patch(facecolor="#0891B2", alpha=0.5, label="Build plate"),
            Patch(facecolor="#059669", label=f"Build dir. (+{bax})"),
        ]
        ax.legend(handles=legend_handles, loc="upper left", fontsize=6.5,
                  frameon=True, framealpha=0.8, edgecolor="#E2E8F0",
                  labelcolor="#334155", borderpad=0.3, handlelength=1.2)

        ax.view_init(elev=22, azim=40)

        plt.tight_layout(pad=0.4)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=150, facecolor=fig.get_facecolor(),
                    bbox_inches="tight")
        plt.close(fig)
        buf.seek(0)
        return buf.read()

    except Exception:
        return None


# ─── Excel ────────────────────────────────────────────────────────────────────

def export_to_excel(
    project_name: str,
    inputs: dict,
    results: dict,
    output_path: str,
    quantity_curve: Optional[list] = None,
) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, GradientFill
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.chart import BarChart, PieChart, LineChart, Reference
        from openpyxl.chart.series import DataPoint

        wb = Workbook()

        # ── Palette ────────────────────────────────────────────────────────────
        NAVY        = "1B2A4A"
        BLUE        = "1E3A8A"
        BLUE_MID    = "2563EB"
        CYAN        = "0891B2"
        CYAN_LIGHT  = "E0F7FA"
        WHITE       = "FFFFFF"
        GREY_LIGHT  = "F8FAFC"
        ALT         = "EFF6FF"
        DARK_ROW    = "1E2A45"
        GOLD        = "F59E0B"

        def _font(bold=False, color=None, size=10, italic=False):
            return Font(name="Calibri", bold=bold, color=color or "374151", size=size, italic=italic)

        _side = Side(style="thin", color="CBD5E1")
        _border = Border(left=_side, right=_side, top=_side, bottom=_side)
        _thick_bottom = Border(left=_side, right=_side, top=_side, bottom=Side(style="medium", color=BLUE_MID))
        _center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _right  = Alignment(horizontal="right",  vertical="center")
        _left   = Alignment(horizontal="left",   vertical="center", indent=1)

        def fill(color): return PatternFill("solid", fgColor=color)

        def hdr_cell(ws, row, col, text, bg=BLUE, fg=WHITE, sz=10, bold=True, align=_center):
            c = ws.cell(row=row, column=col, value=text)
            c.font  = Font(name="Calibri", bold=bold, color=fg, size=sz)
            c.fill  = fill(bg)
            c.border = _border
            c.alignment = align
            return c

        def val_cell(ws, row, col, value, bold=False, color="374151", align=None, fmt=None, bg=None):
            c = ws.cell(row=row, column=col, value=value)
            c.font   = _font(bold=bold, color=color)
            c.border = _border
            c.alignment = align or (_right if isinstance(value, (int, float)) else _left)
            if fmt: c.number_format = fmt
            if bg:  c.fill = fill(bg)
            return c

        def section_header(ws, row, text, ncols, bg=CYAN, fg=WHITE):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
            c = ws.cell(row=row, column=1, value=f"  {text}")
            c.font      = Font(name="Calibri", bold=True, color=fg, size=10)
            c.fill      = fill(bg)
            c.border    = _border
            c.alignment = _left
            for col in range(2, ncols + 1):
                ws.cell(row=row, column=col).fill   = fill(bg)
                ws.cell(row=row, column=col).border = _border
            return row + 1

        def kv_block(ws, start_row, pairs, ncols=2, alt=True):
            r = start_row
            for i, (label, value) in enumerate(pairs):
                bg = ALT if (alt and i % 2 == 0) else GREY_LIGHT
                val_cell(ws, r, 1, label, bg=bg, align=_left)
                val_cell(ws, r, 2, value, bg=bg)
                r += 1
            return r + 1

        total    = results.get("total_cost", 0) or 1
        cpp      = results.get("cost_per_part", 0) or 0
        prod_run = inputs.get("production_run", 1) or 1

        cat_data = [
            ("Material",                results.get("material",         {}).get("cost", 0)),
            ("Build Preparation",       results.get("build_prep",       {}).get("cost", 0)),
            ("Machine Usage",           results.get("machine",          {}).get("cost", 0)),
            ("Consumables",             results.get("consumables",      {}).get("cost", 0)),
            ("Labor",                   results.get("labor",            {}).get("cost", 0)),
            ("Post-Process",            results.get("post_process",     {}).get("cost", 0)),
            ("Heat Treatment",          results.get("metal_am_extensions", {}).get("heat_treatment_cost", 0)),
            ("Part Separation",         results.get("metal_am_extensions", {}).get("part_separation_cost", 0)),
            ("HIP",                     results.get("metal_am_extensions", {}).get("hip_cost", 0)),
            ("Surface Treatment",       results.get("metal_am_extensions", {}).get("surface_treatment_cost", 0)),
            ("NDT / Inspection",        results.get("metal_am_extensions", {}).get("ndt_cost", 0)),
        ]
        active_cats = [(n, v) for n, v in cat_data if v and v > 0]

        machine_name  = inputs.get("machine_name",  inputs.get("machine",  "–"))
        material_name = inputs.get("material_name", inputs.get("material", "–"))

        # ══════════════════════════════════════════════════════════════════════
        # SHEET 1 — Summary Dashboard
        # ══════════════════════════════════════════════════════════════════════
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.sheet_view.showGridLines = False
        ws1.column_dimensions["A"].width = 26
        ws1.column_dimensions["B"].width = 20
        ws1.column_dimensions["C"].width = 20
        ws1.column_dimensions["D"].width = 16
        ws1.column_dimensions["E"].width = 2   # spacer
        ws1.column_dimensions["F"].width = 28
        ws1.column_dimensions["G"].width = 20

        # Title banner (A1:D2)
        ws1.merge_cells("A1:D2")
        banner = ws1["A1"]
        banner.value    = "Metal AM Cost Analysis Report"
        banner.font     = Font(name="Calibri", bold=True, color=WHITE, size=18)
        banner.fill     = fill(NAVY)
        banner.alignment = Alignment(horizontal="left", vertical="center", indent=2)
        ws1.row_dimensions[1].height = 22
        ws1.row_dimensions[2].height = 22
        for col in range(2, 5):
            ws1.cell(row=1, column=col).fill = fill(NAVY)
            ws1.cell(row=2, column=col).fill = fill(NAVY)

        # Sub-info row
        ws1.merge_cells("A3:D3")
        sub = ws1["A3"]
        sub.value     = f"Project: {project_name}  ·  Machine: {machine_name}  ·  Material: {material_name}  ·  {datetime.now().strftime('%B %d, %Y')}"
        sub.font      = Font(name="Calibri", color="B0C4D8", size=9, italic=True)
        sub.fill      = fill(BLUE)
        sub.alignment = Alignment(horizontal="left", vertical="center", indent=2)
        ws1.row_dimensions[3].height = 16
        for col in range(2, 5):
            ws1.cell(row=3, column=col).fill = fill(BLUE)

        # ── KPI Cards (row 5–8, cols A-D) ────────────────────────────────────
        kpis = [
            ("Cost per Part",  _c(cpp),           BLUE_MID),
            ("Total Run Cost", _c(total),          NAVY),
            ("Parts per Build", str(results.get("parts_per_build", "–")), "0F5132"),
            ("# of Builds",    str(results.get("num_builds", "–")),       "6B21A8"),
        ]
        ws1.row_dimensions[5].height = 8   # spacer
        for i, (label, value, bg) in enumerate(kpis, 1):
            c_label = ws1.cell(row=6, column=i, value=label)
            c_label.font      = Font(name="Calibri", bold=True, color=WHITE, size=9)
            c_label.fill      = fill(bg)
            c_label.alignment = Alignment(horizontal="center", vertical="center")
            c_label.border    = _border

            c_val = ws1.cell(row=7, column=i, value=value)
            c_val.font      = Font(name="Calibri", bold=True, color=WHITE, size=14)
            c_val.fill      = fill(bg)
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.border    = _border

            ws1.row_dimensions[6].height = 16
            ws1.row_dimensions[7].height = 28

        # ── Cost Breakdown table (row 10+) ─────────────────────────────────
        r = 10
        r = section_header(ws1, r, "COST BY CATEGORY", 4, bg=NAVY)
        for col, h in enumerate(["Category", "Total Cost", "Cost / Part", "% of Total"], 1):
            hdr_cell(ws1, r, col, h, bg=BLUE_MID, sz=9)
        r += 1

        for i, (cat, cost) in enumerate(active_cats):
            bg = ALT if i % 2 == 0 else GREY_LIGHT
            val_cell(ws1, r, 1, cat,   bg=bg, align=_left)
            val_cell(ws1, r, 2, round(cost, 2), bg=bg, fmt="$#,##0.00")
            val_cell(ws1, r, 3, round(cost / prod_run, 4), bg=bg, fmt="$#,##0.00")
            val_cell(ws1, r, 4, _pct_float(cost, total) / 100, bg=bg, fmt="0.0%")
            r += 1

        # Total row
        for col, (v, fmt) in enumerate([
            ("TOTAL", None), (round(total, 2), "$#,##0.00"),
            (round(cpp, 2), "$#,##0.00"), (1.0, "0%")
        ], 1):
            c = ws1.cell(row=r, column=col, value=v)
            c.font      = Font(name="Calibri", bold=True, color=WHITE, size=10)
            c.fill      = fill(NAVY)
            c.border    = _border
            c.alignment = _right if col > 1 else _left
            if fmt: c.number_format = fmt
        r += 2

        # ── Part & Machine summary (right column F-G) ─────────────────────
        rr = 6
        rr = section_header(ws1, rr, "PART & PRODUCTION", 2, bg=NAVY)
        ws1.cell(row=rr - 1, column=6).fill = fill(NAVY)
        ws1.cell(row=rr - 1, column=7).fill = fill(NAVY)
        # re-do header in F:G columns
        for col in range(6, 8):
            ws1.cell(row=rr - 1, column=col).fill   = fill(NAVY)
            ws1.cell(row=rr - 1, column=col).border = _border

        def kv_right(ws, row, label, value, i):
            bg = ALT if i % 2 == 0 else GREY_LIGHT
            c1 = ws.cell(row=row, column=6, value=label)
            c2 = ws.cell(row=row, column=7, value=value)
            for c in (c1, c2):
                c.fill   = fill(bg)
                c.border = _border
                c.font   = _font(color="374151", size=9)
            c1.alignment = _left
            c2.alignment = _right

        part_kv = [
            ("Machine",       machine_name),
            ("Material",      material_name),
            ("Production Run", f"{prod_run:,} parts"),
            ("Part Volume",   _f(inputs.get("part_volume_cm3"), 2, " cm³")),
            ("Part Height",   _f((inputs.get("part_height_cm") or 0) * 10, 1, " mm")),
            ("Part Mass",     _f(inputs.get("part_mass_kg"), 3, " kg")),
            ("Support Frac.", _f((inputs.get("support_pct_of_mass") or 0) * 100, 1, "%")),
            ("Build Rate",    _f(inputs.get("material_build_rate_cm3_hr"), 1, " cm³/hr")),
            ("Layer Thickness", _f(inputs.get("layer_thickness_um"), 0, " µm")),
        ]
        rr_start = 7
        for i, (lbl, val) in enumerate(part_kv):
            kv_right(ws1, rr_start + i, lbl, val, i)

        # ══════════════════════════════════════════════════════════════════════
        # SHEET 2 — Charts
        # ══════════════════════════════════════════════════════════════════════
        ws_charts = wb.create_sheet("Charts")
        ws_charts.sheet_view.showGridLines = False
        ws_charts.column_dimensions["A"].width = 2

        # Banner
        ws_charts.merge_cells("B1:N2")
        bc = ws_charts["B1"]
        bc.value     = "Visual Analysis"
        bc.font      = Font(name="Calibri", bold=True, color=WHITE, size=16)
        bc.fill      = fill(NAVY)
        bc.alignment = Alignment(horizontal="left", vertical="center", indent=2)
        ws_charts.row_dimensions[1].height = 20
        ws_charts.row_dimensions[2].height = 20
        for col in range(3, 15):
            ws_charts.cell(row=1, column=col).fill = fill(NAVY)
            ws_charts.cell(row=2, column=col).fill = fill(NAVY)

        # Embed matplotlib bar chart
        if active_cats:
            bar_png = _make_bar_chart(
                [n for n, _ in active_cats],
                [v for _, v in active_cats],
                title="Total Cost by Category ($)",
                figsize=(8, 3.8)
            )
            bar_img = XLImage(io.BytesIO(bar_png))
            bar_img.width, bar_img.height = 560, 266
            ws_charts.add_image(bar_img, "B4")

            pie_png = _make_pie_chart(
                [n for n, _ in active_cats],
                [v for _, v in active_cats],
                title="Cost Distribution",
                figsize=(6, 3.8)
            )
            if pie_png:
                pie_img = XLImage(io.BytesIO(pie_png))
                pie_img.width, pie_img.height = 450, 284
                ws_charts.add_image(pie_img, "B24")

        # Quantity curve chart
        if quantity_curve:
            qc_png = _make_line_chart(
                quantity_curve,
                current_qty=prod_run,
                figsize=(10, 3.8)
            )
            if qc_png:
                qc_img = XLImage(io.BytesIO(qc_png))
                qc_img.width, qc_img.height = 660, 252
                ws_charts.add_image(qc_img, "B44")

        # ══════════════════════════════════════════════════════════════════════
        # SHEET 3 — Detailed Cost Breakdown
        # ══════════════════════════════════════════════════════════════════════
        ws3 = wb.create_sheet("Cost Breakdown")
        ws3.sheet_view.showGridLines = False
        ws3.column_dimensions["A"].width = 34
        ws3.column_dimensions["B"].width = 16
        ws3.column_dimensions["C"].width = 10

        ws3.merge_cells("A1:C1")
        t = ws3["A1"]
        t.value, t.font, t.fill, t.alignment = (
            "Detailed Cost Breakdown",
            Font(name="Calibri", bold=True, color=WHITE, size=14),
            fill(NAVY),
            Alignment(horizontal="left", vertical="center", indent=2),
        )
        ws3.row_dimensions[1].height = 28
        for col in range(2, 4): ws3.cell(row=1, column=col).fill = fill(NAVY)
        ws3["A2"].value = project_name
        ws3["A2"].font  = Font(name="Calibri", color="6B7280", size=9, italic=True)
        ws3.row_dimensions[2].height = 14

        r3 = 4

        def bd_section(ws, r, title, rows):
            r = section_header(ws, r, title, 3, bg=BLUE_MID)
            for col, h in enumerate(["Line Item", "Value", "Unit"], 1):
                hdr_cell(ws, r, col, h, bg=CYAN, sz=9)
            r += 1
            for i, (label, value, unit) in enumerate(rows):
                bg = ALT if i % 2 == 0 else GREY_LIGHT
                val_cell(ws, r, 1, f"  {label}", bg=bg, align=_left)
                c = val_cell(ws, r, 2, value, bg=bg)
                if unit == "$" and isinstance(value, (int, float)):
                    c.number_format = "$#,##0.00"
                val_cell(ws, r, 3, unit if unit != "$" else "", bg=bg, align=_center)
                r += 1
            return r + 1

        mat_r  = results.get("material", {})
        mach_r = results.get("machine", {})
        bp_r   = results.get("build_prep", {})
        cons_r = results.get("consumables", {})
        lab_r  = results.get("labor", {})
        pp_r   = results.get("post_process", {})
        ext_r  = results.get("metal_am_extensions", {})

        r3 = bd_section(ws3, r3, "1.  MATERIAL", [
            ("Total Part Material",    round(mat_r.get("total_part_material_kg", 0), 4), "kg"),
            ("Total Support Material", round(mat_r.get("total_support_material_kg", 0), 4), "kg"),
            ("Total Build Material",   round(mat_r.get("total_build_material_kg", 0), 4), "kg"),
            ("Recycled Material",      round(mat_r.get("recycled_material_kg", 0), 4), "kg"),
            ("Material Waste",         round(mat_r.get("material_waste_kg", 0), 4), "kg"),
            ("Net Material Required",  round(mat_r.get("material_required_kg", 0), 4), "kg"),
            ("Material Cost",          round(mat_r.get("cost", 0), 2), "$"),
        ])
        r3 = bd_section(ws3, r3, "2.  BUILD PREPARATION", [
            ("Total Prep Time",  round(bp_r.get("total_time_hr", 0), 2), "hr"),
            ("Prep Cost",        round(bp_r.get("cost", 0), 2), "$"),
        ])
        r3 = bd_section(ws3, r3, "3.  MACHINE USAGE", [
            ("Machine Cost/Hour",   round(mach_r.get("machine_cost_per_hour", 0), 2), "$/hr"),
            ("Warmup Time",         round(mach_r.get("total_warmup_hr", 0), 2), "hr"),
            ("Laser Scan Time",     round(mach_r.get("total_scan_hr", 0), 2), "hr"),
            ("Recoating Time",      round(mach_r.get("total_recoating_hr", 0), 2), "hr"),
            ("Total Print Time",    round(mach_r.get("total_print_hr", 0), 2), "hr"),
            ("Cooldown Time",       round(mach_r.get("total_cooldown_hr", 0), 2), "hr"),
            ("Build Exchange Time", round(mach_r.get("total_build_exchange_hr", 0), 2), "hr"),
            ("Total Machine Time",  round(mach_r.get("total_machine_time_hr", 0), 2), "hr"),
            ("Machine Cost",        round(mach_r.get("cost", 0), 2), "$"),
        ])
        r3 = bd_section(ws3, r3, "4.  CONSUMABLES", [
            ("Inert Gas Cost",        round(cons_r.get("inert_gas_cost", 0), 2), "$"),
            ("Per-Build Consumables", round(cons_r.get("per_build_consumables", 0), 2), "$"),
            ("Total Consumables",     round(cons_r.get("cost", 0), 2), "$"),
        ])
        r3 = bd_section(ws3, r3, "5.  LABOR", [
            ("Labor During Build",   round(lab_r.get("labor_during_build_hr", 0), 2), "hr"),
            ("Build Exchange Labor", round(lab_r.get("labor_build_exchange_hr", 0), 2), "hr"),
            ("Total Labor",          round(lab_r.get("total_labor_hr", 0), 2), "hr"),
            ("Labor Cost",           round(lab_r.get("cost", 0), 2), "$"),
        ])
        r3 = bd_section(ws3, r3, "6.  POST-PROCESS — Support Removal", [
            ("Support Removal Time", round(pp_r.get("support_removal_time_hr", 0), 2), "hr"),
            ("Post-Process Cost",    round(pp_r.get("cost", 0), 2), "$"),
        ])
        r3 = bd_section(ws3, r3, "7.  METAL AM POST-PROCESSING", [
            ("Heat Treatment",    round(ext_r.get("heat_treatment_cost", 0), 2), "$"),
            ("Part Separation",   round(ext_r.get("part_separation_cost", 0), 2), "$"),
            ("HIP",               round(ext_r.get("hip_cost", 0), 2), "$"),
            ("Surface Treatment", round(ext_r.get("surface_treatment_cost", 0), 2), "$"),
            ("NDT / Inspection",  round(ext_r.get("ndt_cost", 0), 2), "$"),
        ])

        # ══════════════════════════════════════════════════════════════════════
        # SHEET 4 — Machine & Process
        # ══════════════════════════════════════════════════════════════════════
        ws4 = wb.create_sheet("Machine & Process")
        ws4.sheet_view.showGridLines = False
        ws4.column_dimensions["A"].width = 30
        ws4.column_dimensions["B"].width = 26

        ws4.merge_cells("A1:B1")
        h4 = ws4["A1"]
        h4.value, h4.font, h4.fill, h4.alignment = (
            "Machine & Process Configuration",
            Font(name="Calibri", bold=True, color=WHITE, size=14),
            fill(NAVY), Alignment(horizontal="left", vertical="center", indent=2),
        )
        ws4.row_dimensions[1].height = 28
        ws4["B1"].fill = fill(NAVY)

        r4 = 3
        r4 = section_header(ws4, r4, "MACHINE SPECIFICATIONS", 2, bg=BLUE_MID)
        r4 = kv_block(ws4, r4, [
            ("Machine Name",        inputs.get("machine_name", "–")),
            ("Manufacturer",        inputs.get("machine_manufacturer", "–")),
            ("Build Plate",         inputs.get("machine_build_plate", "–")),
            ("Build Height",        _f(inputs.get("machine_build_height_mm"), 0, " mm")),
            ("Build Area",          _f(inputs.get("machine_build_area_cm2"), 1, " cm²")),
            ("Machine Purchase Cost", _c(inputs.get("machine_cost_usd", 0))),
            ("Number of Lasers",    str(inputs.get("machine_num_lasers", 1))),
            ("Packing Policy",      inputs.get("machine_packing_policy", "2D")),
            ("Packing Fraction",    _f((inputs.get("machine_packing_fraction") or 0.9) * 100, 0, "%")),
            ("Machine Uptime",      _f((inputs.get("machine_uptime") or 0.8) * 100, 0, "%")),
            ("Powder Recycling",    _f((inputs.get("machine_recycling_fraction") or 0.9) * 100, 0, "%")),
        ])
        r4 = section_header(ws4, r4, "PROCESS PARAMETERS", 2, bg=BLUE_MID)
        r4 = kv_block(ws4, r4, [
            ("Build Direction",         (inputs.get("build_direction") or "–").upper()),
            ("Elevation above Plate",   _f(inputs.get("elevation_mm", 0), 0, " mm")),
            ("Material Build Rate",     _f(inputs.get("material_build_rate_cm3_hr"), 1, " cm³/hr")),
            ("Layer Thickness",         _f(inputs.get("layer_thickness_um"), 0, " µm")),
            ("Laser Strategy",          inputs.get("laser_strategy", "single")),
            ("Lasers per Part",         str(inputs.get("lasers_per_part", 1))),
            ("Recoating Efficiency",    _f((inputs.get("recoating_efficiency") or 0.75) * 100, 0, "%")),
            ("Parts per Build Override", str(inputs.get("parts_per_build_override") or "Auto")),
        ])
        r4 = section_header(ws4, r4, "POST-PROCESSING INPUTS", 2, bg=BLUE_MID)
        r4 = kv_block(ws4, r4, [
            ("Heat Treatment",       "Enabled" if inputs.get("heat_treatment_enabled") else "Disabled"),
            ("  → Cost per Build",   _c(inputs.get("heat_treatment_cost_per_build", 0))),
            ("Part Separation",      "Enabled" if inputs.get("part_separation_enabled") else "Disabled"),
            ("  → Cost per Part",    _c(inputs.get("part_separation_cost_per_part", 0))),
            ("HIP",                  "Enabled" if inputs.get("hip_enabled") else "Disabled"),
            ("  → Cost per Part",    _c(inputs.get("hip_cost_per_part", 0))),
            ("Surface Treatment",    "Enabled" if inputs.get("surface_treatment_enabled") else "Disabled"),
            ("  → Cost per Part",    _c(inputs.get("surface_treatment_cost_per_part", 0))),
            ("NDT / Inspection",     "Enabled" if inputs.get("ndt_enabled") else "Disabled"),
            ("  → Cost per Part",    _c(inputs.get("ndt_cost_per_part", 0))),
        ])

        # ══════════════════════════════════════════════════════════════════════
        # SHEET 5 — Material & Part
        # ══════════════════════════════════════════════════════════════════════
        ws5 = wb.create_sheet("Material & Part")
        ws5.sheet_view.showGridLines = False
        ws5.column_dimensions["A"].width = 28
        ws5.column_dimensions["B"].width = 26

        ws5.merge_cells("A1:B1")
        h5 = ws5["A1"]
        h5.value, h5.font, h5.fill, h5.alignment = (
            "Material & Part Specification",
            Font(name="Calibri", bold=True, color=WHITE, size=14),
            fill(NAVY), Alignment(horizontal="left", vertical="center", indent=2),
        )
        ws5.row_dimensions[1].height = 28
        ws5["B1"].fill = fill(NAVY)

        r5 = 3
        r5 = section_header(ws5, r5, "MATERIAL", 2, bg=BLUE_MID)
        r5 = kv_block(ws5, r5, [
            ("Name",             inputs.get("material_name", "–")),
            ("Alloy Designation", inputs.get("material_alloy", "–")),
            ("Category",         inputs.get("material_category", "–")),
            ("Cost per kg",      _c(inputs.get("material_cost_per_kg", 0))),
            ("Density",          _f(inputs.get("material_density_g_cm3"), 2, " g/cm³")),
            ("Layer Thickness",  _f(inputs.get("layer_thickness_um"), 0, " µm")),
            ("Build Rate",       _f(inputs.get("material_build_rate_cm3_hr"), 1, " cm³/hr")),
        ])
        r5 = section_header(ws5, r5, "PART GEOMETRY", 2, bg=BLUE_MID)
        r5 = kv_block(ws5, r5, [
            ("Part File",          inputs.get("part_file_name", "–")),
            ("Volume",             _f(inputs.get("part_volume_cm3"), 3, " cm³")),
            ("Height (build dir.)", _f((inputs.get("part_height_cm") or 0) * 10, 2, " mm")),
            ("Projected Area",     _f(inputs.get("part_projected_area_cm2"), 2, " cm²")),
            ("Mass (computed)",    _f(inputs.get("part_mass_kg"), 4, " kg")),
            ("Support Fraction",   _f((inputs.get("support_pct_of_mass") or 0) * 100, 1, "%")),
            ("Elevation above Plate", _f(inputs.get("elevation_mm", 0), 0, " mm")),
        ])
        r5 = section_header(ws5, r5, "PRODUCTION", 2, bg=BLUE_MID)
        r5 = kv_block(ws5, r5, [
            ("Production Run",    f"{prod_run:,} parts"),
            ("Parts per Build",   str(results.get("parts_per_build", "–"))),
            ("Number of Builds",  str(results.get("num_builds", "–"))),
            ("Total Machine Time", _f(mach_r.get("total_machine_time_hr", 0), 1, " hr")),
        ])

        # ══════════════════════════════════════════════════════════════════════
        # SHEET 6 — Quantity Curve
        # ══════════════════════════════════════════════════════════════════════
        if quantity_curve:
            ws6 = wb.create_sheet("Quantity Curve")
            ws6.sheet_view.showGridLines = False
            for col_letter, w in zip("ABCDEFGHIJK", [8,14,16,14,14,14,14,14,14,10,8]):
                ws6.column_dimensions[col_letter].width = w

            ws6.merge_cells("A1:K1")
            h6 = ws6["A1"]
            h6.value, h6.font, h6.fill, h6.alignment = (
                "Cost vs Quantity Data",
                Font(name="Calibri", bold=True, color=WHITE, size=14),
                fill(NAVY), Alignment(horizontal="left", vertical="center", indent=2),
            )
            ws6.row_dimensions[1].height = 28
            for col in range(2, 12): ws6.cell(row=1, column=col).fill = fill(NAVY)

            headers = [
                "Qty", "Cost/Part ($)", "Total Cost ($)", "Material/Part",
                "Machine/Part", "Labor/Part", "Consumables/Part",
                "Build Prep/Part", "Post-Process/Part", "Parts/Build", "# Builds",
            ]
            r6 = 3
            for c, h in enumerate(headers, 1):
                hdr_cell(ws6, r6, c, h, bg=BLUE_MID, sz=9)

            for i, pt in enumerate(quantity_curve):
                r6 += 1
                vals = [
                    pt.get("quantity"), pt.get("cost_per_part"), pt.get("total_cost"),
                    pt.get("material_cost"), pt.get("machine_cost"), pt.get("labor_cost"),
                    pt.get("consumables_cost"), pt.get("build_prep_cost"),
                    pt.get("post_process_cost"), pt.get("parts_per_build"), pt.get("num_builds"),
                ]
                bg = ALT if i % 2 == 0 else GREY_LIGHT
                for c, v in enumerate(vals, 1):
                    cell = ws6.cell(row=r6, column=c, value=round(v, 4) if isinstance(v, float) else v)
                    cell.fill   = fill(bg)
                    cell.border = _border
                    cell.alignment = _right
                    cell.font   = _font(size=9)
                    if 2 <= c <= 9:
                        cell.number_format = "$#,##0.00"

            # Freeze header row
            ws6.freeze_panes = "A4"

        wb.save(output_path)
        return output_path
    except Exception as e:
        raise Exception(f"Excel export failed: {str(e)}")


# ─── PDF ──────────────────────────────────────────────────────────────────────

def export_to_pdf(
    project_name: str,
    inputs: dict,
    results: dict,
    output_path: str,
    quantity_curve: Optional[list] = None,
) -> str:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph,
            Spacer, KeepTogether, Image as RLImage,
        )
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

        PAGE_W, PAGE_H = A4
        MARGIN  = 1.8 * cm
        COL_W   = PAGE_W - 2 * MARGIN

        # ── Palette ───────────────────────────────────────────────────────────
        NAVY     = colors.HexColor("#1B2A4A")
        BLUE_MID = colors.HexColor("#2563EB")
        ALT      = colors.HexColor("#EFF6FF")
        WHITE    = colors.white
        DARK     = colors.HexColor("#111827")
        MUTED    = colors.HexColor("#64748B")
        GOLD     = colors.HexColor("#F59E0B")
        SUCCESS  = colors.HexColor("#059669")
        PURPLE   = colors.HexColor("#7C3AED")
        DIVIDER  = colors.HexColor("#CBD5E1")

        # ── Style factory ─────────────────────────────────────────────────────
        def _s(name, **kw):
            base = dict(fontName="Helvetica", fontSize=9, leading=12, textColor=DARK)
            base.update(kw)
            return ParagraphStyle(name, **base)

        s_title   = _s("Ti", fontSize=20, textColor=WHITE, fontName="Helvetica-Bold",
                        leading=24)
        s_sub     = _s("Su", fontSize=8.5, textColor=colors.HexColor("#B0C4D8"),
                        leading=12)
        s_label   = _s("Lb", fontSize=8, textColor=MUTED, leading=11)
        s_val     = _s("Vl", fontSize=8.5, textColor=DARK, fontName="Helvetica-Bold",
                        leading=11)
        s_kpi_lbl = _s("KL", fontSize=7.5, textColor=WHITE, fontName="Helvetica-Bold",
                        alignment=TA_CENTER, leading=10)
        s_kpi_val = _s("KV", fontSize=15, textColor=WHITE, fontName="Helvetica-Bold",
                        alignment=TA_CENTER, leading=19)
        s_tbl_hdr = _s("TH", fontSize=8, textColor=WHITE, fontName="Helvetica-Bold",
                        alignment=TA_CENTER, leading=10)
        s_ack     = _s("Ak", fontSize=6.5, textColor=MUTED, alignment=TA_CENTER,
                        leading=9)

        # ── Table helpers ─────────────────────────────────────────────────────
        _BORDER = DIVIDER

        def _base_ts(nrows=0, hdr_rows=1):
            ts = TableStyle([
                ('BACKGROUND',   (0, 0), (-1, hdr_rows - 1), BLUE_MID),
                ('TEXTCOLOR',    (0, 0), (-1, hdr_rows - 1), WHITE),
                ('FONTNAME',     (0, 0), (-1, hdr_rows - 1), 'Helvetica-Bold'),
                ('FONTSIZE',     (0, 0), (-1, -1), 8),
                ('GRID',         (0, 0), (-1, -1), 0.25, _BORDER),
                ('LEFTPADDING',  (0, 0), (-1, -1), 7),
                ('RIGHTPADDING', (0, 0), (-1, -1), 7),
                ('TOPPADDING',   (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING',(0, 0), (-1, -1), 4),
                ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN',        (1, 0), (-1, -1), 'RIGHT'),
                ('ALIGN',        (0, 0), (0, -1),  'LEFT'),
            ])
            for i in range(hdr_rows, nrows, 2):
                ts.add('BACKGROUND', (0, i), (-1, i), ALT)
            return ts

        def _kv_tbl(rows, cw=None):
            """Two-column key/value table with muted labels and bold values."""
            cw = cw or [COL_W * 0.44, COL_W * 0.56]
            data = [[Paragraph(lbl, s_label), Paragraph(str(val), s_val)]
                    for lbl, val in rows]
            tbl = Table(data, colWidths=cw)
            ts = TableStyle([
                ('GRID',          (0, 0), (-1, -1), 0.25, _BORDER),
                ('LEFTPADDING',   (0, 0), (-1, -1), 8),
                ('RIGHTPADDING',  (0, 0), (-1, -1), 8),
                ('TOPPADDING',    (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN',         (1, 0), (1, -1),  'RIGHT'),
                ('ALIGN',         (0, 0), (0, -1),  'LEFT'),
            ])
            for i in range(0, len(rows), 2):
                ts.add('BACKGROUND', (0, i), (-1, i), ALT)
            tbl.setStyle(ts)
            return tbl

        def _sec_hdr(text):
            """Blue section header with gold left accent bar."""
            data = [["", Paragraph(text.upper(), _s("SH", fontSize=8, textColor=WHITE,
                                                     fontName="Helvetica-Bold", leading=10))]]
            tbl = Table(data, colWidths=[4, COL_W - 4])
            tbl.setStyle(TableStyle([
                ('BACKGROUND',   (0, 0), (0, 0), GOLD),
                ('BACKGROUND',   (1, 0), (1, 0), BLUE_MID),
                ('LEFTPADDING',  (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ('TOPPADDING',   (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING',(0, 0), (-1, -1), 5),
                ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
            ]))
            return tbl

        def _chart_img(png_bytes, target_w, fw=None, fh=None):
            """Place chart PNG maintaining its original aspect ratio.
            Reads actual pixel dimensions from the PNG header so tight_layout
            adjustments don't cause stretching."""
            import struct
            px_w = struct.unpack(">I", png_bytes[16:20])[0]
            px_h = struct.unpack(">I", png_bytes[20:24])[0]
            target_h = target_w * (px_h / px_w) if px_w else target_w * 0.5
            return RLImage(io.BytesIO(png_bytes), width=target_w, height=target_h)

        # ── Per-page footer (page number + acknowledgment) ────────────────────
        def _footer(canvas, doc):
            canvas.saveState()
            fy = 0.85 * cm
            canvas.setStrokeColor(DIVIDER)
            canvas.setLineWidth(0.3)
            canvas.line(MARGIN, fy + 9, PAGE_W - MARGIN, fy + 9)
            canvas.setFont("Helvetica", 6.5)
            canvas.setFillColor(MUTED)
            canvas.drawString(
                MARGIN, fy,
                "Tool developed by Chinmay Phutela  ·  Cost model methodology inspired by MIT  ·  Metal AM Cost Analyzer"
            )
            canvas.drawRightString(PAGE_W - MARGIN, fy, f"Page {doc.page}")
            canvas.restoreState()

        doc = SimpleDocTemplate(
            output_path, pagesize=A4,
            leftMargin=MARGIN, rightMargin=MARGIN,
            topMargin=MARGIN, bottomMargin=1.9 * cm,
        )

        # ── Common data ───────────────────────────────────────────────────────
        story = []
        total         = results.get("total_cost", 0) or 1
        cpp           = results.get("cost_per_part", 0) or 0
        prod_run      = inputs.get("production_run", 1) or 1
        machine_name  = inputs.get("machine_name",  inputs.get("machine",  "–"))
        material_name = inputs.get("material_name", inputs.get("material", "–"))

        categories = [
            ("Material",         results.get("material",    {}).get("cost", 0)),
            ("Build Prep",       results.get("build_prep",  {}).get("cost", 0)),
            ("Machine Usage",    results.get("machine",     {}).get("cost", 0)),
            ("Consumables",      results.get("consumables", {}).get("cost", 0)),
            ("Labor",            results.get("labor",       {}).get("cost", 0)),
            ("Post-Process",     results.get("post_process",{}).get("cost", 0)),
            ("Heat Treatment",   results.get("metal_am_extensions", {}).get("heat_treatment_cost", 0)),
            ("Part Separation",  results.get("metal_am_extensions", {}).get("part_separation_cost", 0)),
            ("HIP",              results.get("metal_am_extensions", {}).get("hip_cost", 0)),
            ("Surface Treatment",results.get("metal_am_extensions", {}).get("surface_treatment_cost", 0)),
            ("NDT / Inspection", results.get("metal_am_extensions", {}).get("ndt_cost", 0)),
        ]
        active_cats = [(n, v) for n, v in categories if v and v > 0]

        # ═════════════════════════════════════════════════════════════════════
        # COVER BANNER
        # ═════════════════════════════════════════════════════════════════════
        cover_data = [[
            Paragraph("Metal AM Cost Analysis Report", s_title),
            Paragraph(
                f"<b>Project:</b> {project_name}<br/>"
                f"<b>Machine:</b> {machine_name}<br/>"
                f"<b>Material:</b> {material_name}<br/>"
                f"<b>Date:</b> {datetime.now().strftime('%B %d, %Y')}",
                s_sub
            ),
        ]]
        cover = Table(cover_data, colWidths=[COL_W * 0.55, COL_W * 0.45])
        cover.setStyle(TableStyle([
            ('BACKGROUND',   (0, 0), (-1, -1), NAVY),
            ('LEFTPADDING',  (0, 0), (-1, -1), 14),
            ('RIGHTPADDING', (0, 0), (-1, -1), 14),
            ('TOPPADDING',   (0, 0), (-1, -1), 18),
            ('BOTTOMPADDING',(0, 0), (-1, -1), 18),
            ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
            ('LINEBELOW',    (0, 0), (-1, 0), 3, GOLD),
        ]))
        story.append(cover)
        story.append(Spacer(1, 10))

        # ═════════════════════════════════════════════════════════════════════
        # KPI CARDS
        # ═════════════════════════════════════════════════════════════════════
        kpis = [
            ("Cost per Part",   _c(cpp),   BLUE_MID),
            ("Total Run Cost",  _c(total),  NAVY),
            ("Parts per Build", str(results.get("parts_per_build", "–")), SUCCESS),
            ("# of Builds",     str(results.get("num_builds", "–")),      PURPLE),
        ]
        kpi_data = [
            [Paragraph(lbl, s_kpi_lbl) for lbl, _, _ in kpis],
            [Paragraph(val, s_kpi_val) for _, val, _ in kpis],
        ]
        kpi_tbl = Table(kpi_data, colWidths=[COL_W / 4] * 4, rowHeights=[18, 32])
        kpi_ts = TableStyle([
            ('LEFTPADDING',  (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING',   (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING',(0, 0), (-1, -1), 5),
            ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN',        (0, 0), (-1, -1), 'CENTER'),
            ('LINEBELOW',    (0, 1), (-1, 1), 0.5, _BORDER),
        ])
        for i, (_, _, c) in enumerate(kpis):
            kpi_ts.add('BACKGROUND', (i, 0), (i, 1), c)
        kpi_tbl.setStyle(kpi_ts)
        story.append(kpi_tbl)
        story.append(Spacer(1, 14))

        # ═════════════════════════════════════════════════════════════════════
        # STL GEOMETRY SNAPSHOT (optional — only if file is still on disk)
        # ═════════════════════════════════════════════════════════════════════
        stl_path = inputs.get("part_file_path") or inputs.get("file_path")
        build_dir = inputs.get("build_direction", "z") or "z"
        if stl_path:
            STL_FS = (9.0, 4.2)   # figsize (inches)
            stl_png = _render_stl_snapshot(stl_path, build_direction=build_dir,
                                           overhang_angle=45.0, figsize=STL_FS)
            if stl_png:
                story.append(_sec_hdr("Part Geometry — 3D View · Overhangs · Bounding Box"))
                story.append(Spacer(1, 6))

                stl_im = _chart_img(stl_png, COL_W - 8, *STL_FS)
                stl_tbl = Table([[stl_im]], colWidths=[COL_W])
                stl_tbl.setStyle(TableStyle([
                    ('BACKGROUND',   (0, 0), (-1, -1), WHITE),
                    ('ALIGN',        (0, 0), (-1, -1), 'CENTER'),
                    ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
                    ('LEFTPADDING',  (0, 0), (-1, -1), 2),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 2),
                    ('TOPPADDING',   (0, 0), (-1, -1), 2),
                    ('BOTTOMPADDING',(0, 0), (-1, -1), 2),
                    ('BOX',          (0, 0), (-1, -1), 0.5, _BORDER),
                ]))
                story.append(stl_tbl)

                # Small geometry stats row below the image
                bb = inputs.get("bounding_box_mm", {})
                geom_rows = [
                    Paragraph(f"<b>Volume:</b> {_f(inputs.get('part_volume_cm3'), 3)} cm³", s_label),
                    Paragraph(f"<b>Mass:</b> {_f(inputs.get('part_mass_kg'), 4)} kg", s_label),
                    Paragraph(f"<b>Height:</b> {_f((inputs.get('part_height_cm') or 0)*10, 1)} mm", s_label),
                    Paragraph(f"<b>Support est.:</b> {_f((inputs.get('support_pct_of_mass') or 0)*100, 1)}%", s_label),
                ]
                gr_tbl = Table([geom_rows], colWidths=[COL_W / 4] * 4)
                gr_tbl.setStyle(TableStyle([
                    ('BACKGROUND',   (0, 0), (-1, -1), ALT),
                    ('ALIGN',        (0, 0), (-1, -1), 'CENTER'),
                    ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
                    ('TOPPADDING',   (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING',(0, 0), (-1, -1), 5),
                    ('LEFTPADDING',  (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                    ('BOX',          (0, 0), (-1, -1), 0.5, _BORDER),
                    ('LINEABOVE',    (0, 0), (-1, 0), 0.25, _BORDER),
                ]))
                story.append(gr_tbl)
                story.append(Spacer(1, 12))

        # ═════════════════════════════════════════════════════════════════════
        # CHARTS  — aspect-ratio-correct placement
        # ═════════════════════════════════════════════════════════════════════
        # figsize chosen so both charts share the same pixel height at their
        # respective target widths → no vertical stretching.
        BAR_FS = (7.2, 3.4)   # width, height in inches
        PIE_FS = (4.8, 3.4)
        LC_FS  = (10.0, 3.4)

        # Both charts rendered at H=136 pts; derive widths from aspect ratio
        H_CHART = 136
        W_BAR   = H_CHART * (BAR_FS[0] / BAR_FS[1])   # ≈ 288 pts
        W_PIE   = H_CHART * (PIE_FS[0] / PIE_FS[1])   # ≈ 192 pts
        # column widths fill COL_W exactly
        W_BAR_COL = COL_W - W_PIE - 6
        W_PIE_COL = W_PIE + 6

        if active_cats:
            story.append(_sec_hdr("Cost Analysis Charts"))
            story.append(Spacer(1, 6))

            bar_png = _make_bar_chart(
                [n for n, _ in active_cats], [v for _, v in active_cats],
                title="Cost by Category ($)", figsize=BAR_FS,
            )
            pie_png = _make_pie_chart(
                [n for n, _ in active_cats], [v for _, v in active_cats],
                title="Cost Distribution", figsize=PIE_FS,
            )
            bar_im = _chart_img(bar_png, W_BAR_COL - 6, *BAR_FS)

            if pie_png:
                pie_im = _chart_img(pie_png, W_PIE, *PIE_FS)
                ct = Table([[bar_im, pie_im]], colWidths=[W_BAR_COL, W_PIE_COL])
            else:
                bar_im = _chart_img(bar_png, COL_W - 8, *BAR_FS)
                ct = Table([[bar_im]], colWidths=[COL_W])

            ct.setStyle(TableStyle([
                ('BACKGROUND',   (0, 0), (-1, -1), WHITE),
                ('ALIGN',        (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING',  (0, 0), (-1, -1), 2),
                ('RIGHTPADDING', (0, 0), (-1, -1), 2),
                ('TOPPADDING',   (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING',(0, 0), (-1, -1), 2),
                ('BOX',          (0, 0), (-1, -1), 0.5, _BORDER),
            ]))
            story.append(ct)
            story.append(Spacer(1, 12))

        if quantity_curve:
            lc_png = _make_line_chart(quantity_curve, current_qty=prod_run, figsize=LC_FS)
            if lc_png:
                story.append(_sec_hdr("Cost per Part vs Production Quantity"))
                story.append(Spacer(1, 6))
                lc_im = _chart_img(lc_png, COL_W - 8, *LC_FS)
                lct = Table([[lc_im]], colWidths=[COL_W])
                lct.setStyle(TableStyle([
                    ('BACKGROUND',   (0, 0), (-1, -1), WHITE),
                    ('ALIGN',        (0, 0), (-1, -1), 'CENTER'),
                    ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
                    ('LEFTPADDING',  (0, 0), (-1, -1), 2),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 2),
                    ('TOPPADDING',   (0, 0), (-1, -1), 2),
                    ('BOTTOMPADDING',(0, 0), (-1, -1), 2),
                    ('BOX',          (0, 0), (-1, -1), 0.5, _BORDER),
                ]))
                story.append(lct)
                story.append(Spacer(1, 12))

        # ═════════════════════════════════════════════════════════════════════
        # COST BREAKDOWN TABLE
        # ═════════════════════════════════════════════════════════════════════
        story.append(_sec_hdr("Cost Breakdown by Category"))
        story.append(Spacer(1, 6))

        bk_hdrs = [
            Paragraph(h, _s("BH", fontSize=8, textColor=WHITE,
                             fontName="Helvetica-Bold",
                             alignment=TA_LEFT if i == 0 else TA_RIGHT))
            for i, h in enumerate(["Category", "Total Cost", "Cost / Part", "% of Total"])
        ]
        bk_data = [bk_hdrs]
        for cat, cost in active_cats:
            bk_data.append([cat, _c(cost), _c(cost / prod_run), _pct(cost, total)])
        bk_data.append(["TOTAL", _c(total), _c(cpp), "100%"])

        bk_tbl = Table(bk_data, colWidths=[COL_W*0.36, COL_W*0.22, COL_W*0.22, COL_W*0.20])
        bk_ts = _base_ts(nrows=len(bk_data))
        last = len(bk_data) - 1
        bk_ts.add('BACKGROUND', (0, last), (-1, last), NAVY)
        bk_ts.add('TEXTCOLOR',  (0, last), (-1, last), WHITE)
        bk_ts.add('FONTNAME',   (0, last), (-1, last), 'Helvetica-Bold')
        bk_tbl.setStyle(bk_ts)
        story.append(KeepTogether([bk_tbl]))
        story.append(Spacer(1, 12))

        # ═════════════════════════════════════════════════════════════════════
        # MACHINE TIME BREAKDOWN
        # ═════════════════════════════════════════════════════════════════════
        mach_r = results.get("machine", {})
        story.append(_sec_hdr("Machine Time Breakdown"))
        story.append(Spacer(1, 6))

        time_hdrs = [Paragraph(h, s_tbl_hdr)
                     for h in ["Phase", "Time (hr)", "Notes"]]
        time_rows = [
            ("Warmup",         _f(mach_r.get("total_warmup_hr", 0), 2),
                                f"{results.get('num_builds', 0)} builds × warmup time"),
            ("Laser Scan",     _f(mach_r.get("total_scan_hr", 0), 2),
                                "Laser deposition time"),
            ("Recoating",      _f(mach_r.get("total_recoating_hr", 0), 2),
                                "Powder spreading overhead"),
            ("Cooldown",       _f(mach_r.get("total_cooldown_hr", 0), 2),
                                f"{results.get('num_builds', 0)} builds × cooldown time"),
            ("Build Exchange", _f(mach_r.get("total_build_exchange_hr", 0), 2),
                                "Setup and plate removal"),
            ("TOTAL",          _f(mach_r.get("total_machine_time_hr", 0), 2), ""),
        ]
        t_data = [time_hdrs] + [list(r) for r in time_rows]
        t_tbl = Table(t_data, colWidths=[COL_W*0.22, COL_W*0.16, COL_W*0.62])
        t_ts = _base_ts(nrows=len(t_data))
        tl = len(t_data) - 1
        t_ts.add('BACKGROUND', (0, tl), (-1, tl), NAVY)
        t_ts.add('TEXTCOLOR',  (0, tl), (-1, tl), WHITE)
        t_ts.add('FONTNAME',   (0, tl), (-1, tl), 'Helvetica-Bold')
        t_tbl.setStyle(t_ts)
        story.append(KeepTogether([t_tbl]))
        story.append(Spacer(1, 12))

        # ═════════════════════════════════════════════════════════════════════
        # MATERIAL & PART DETAILS (two-column)
        # ═════════════════════════════════════════════════════════════════════
        mat_r = results.get("material", {})
        story.append(_sec_hdr("Material & Part Details"))
        story.append(Spacer(1, 6))

        half = (COL_W - 8) / 2
        mat_rows = [
            ("Total Part Material",    _f(mat_r.get("total_part_material_kg", 0), 3, " kg")),
            ("Total Support Material", _f(mat_r.get("total_support_material_kg", 0), 3, " kg")),
            ("Net Material Required",  _f(mat_r.get("material_required_kg", 0), 3, " kg")),
            ("Powder Recycled",        _f(mat_r.get("recycled_material_kg", 0), 3, " kg")),
            ("Material Waste",         _f(mat_r.get("material_waste_kg", 0), 3, " kg")),
            ("Material Cost",          _c(mat_r.get("cost", 0))),
        ]
        part_rows = [
            ("Part File",           inputs.get("part_file_name", "–")),
            ("Volume",              _f(inputs.get("part_volume_cm3"), 3, " cm³")),
            ("Height (build dir.)", _f((inputs.get("part_height_cm") or 0) * 10, 2, " mm")),
            ("Mass",                _f(inputs.get("part_mass_kg"), 4, " kg")),
            ("Support Fraction",    _f((inputs.get("support_pct_of_mass") or 0) * 100, 1, "%")),
            ("Production Run",      f"{prod_run:,} parts"),
        ]
        two_col = Table(
            [[_kv_tbl(mat_rows,  [half * 0.48, half * 0.52]),
              _kv_tbl(part_rows, [half * 0.48, half * 0.52])]],
            colWidths=[half + 4, half + 4],
        )
        two_col.setStyle(TableStyle([
            ('VALIGN',       (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING',  (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, -1),  4),
            ('RIGHTPADDING', (1, 0), (1, -1),  0),
        ]))
        story.append(KeepTogether([two_col]))
        story.append(Spacer(1, 12))

        # ═════════════════════════════════════════════════════════════════════
        # MACHINE & PROCESS CONFIGURATION
        # ═════════════════════════════════════════════════════════════════════
        story.append(_sec_hdr("Machine & Process Configuration"))
        story.append(Spacer(1, 6))

        machine_rows = [
            ("Machine",              inputs.get("machine_name", "–")),
            ("Manufacturer",         inputs.get("machine_manufacturer", "–")),
            ("Build Plate",          inputs.get("machine_build_plate", "–")),
            ("Build Height",         _f(inputs.get("machine_build_height_mm"), 0, " mm")),
            ("Build Area",           _f(inputs.get("machine_build_area_cm2"), 1, " cm²")),
            ("Machine Purchase Cost",_c(inputs.get("machine_cost_usd", 0))),
            ("Number of Lasers",     str(inputs.get("machine_num_lasers", 1))),
            ("Laser Strategy",       inputs.get("laser_strategy", "single")),
            ("Material",             inputs.get("material_name", "–")),
            ("Material Cost / kg",   _c(inputs.get("material_cost_per_kg", 0))),
            ("Build Rate",           _f(inputs.get("material_build_rate_cm3_hr"), 1, " cm³/hr")),
            ("Layer Thickness",      _f(inputs.get("layer_thickness_um"), 0, " µm")),
            ("Packing Policy",       inputs.get("machine_packing_policy", "2D")),
            ("Packing Fraction",     _f((inputs.get("machine_packing_fraction") or 0.9) * 100, 0, "%")),
            ("Recoating Efficiency", _f((inputs.get("recoating_efficiency") or 0.75) * 100, 0, "%")),
        ]
        story.append(KeepTogether([_kv_tbl(machine_rows)]))
        story.append(Spacer(1, 12))

        # ═════════════════════════════════════════════════════════════════════
        # BUILD
        # ═════════════════════════════════════════════════════════════════════
        doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
        return output_path
    except Exception as e:
        raise Exception(f"PDF export failed: {str(e)}")
